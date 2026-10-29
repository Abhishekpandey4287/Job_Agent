"""Dedicated workbook for product-company Java openings. Existing agent exports are not used."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from product_java_scan.catalog import Company
from product_java_scan.models import CompanyResult, FoundJob
from product_java_scan.scoring import experience_acceptable
from product_java_scan.textutil import posted_within_week
from product_java_scan.urls import canonical_url

VACANCY_HEADERS = [
    "Company Name",
    "Job Title",
    "Location",
    "Experience",
    "Required Skills",
    "Job URL",
    "Posting Date",
    "Match Score",
    "Role Match",
    "Verified Via",
    "Location Check",
    "First Seen",
    "Last Seen",
    "Record Status",
    "Actively Hiring Java/Spring",
]
COMPANY_HEADERS = [
    "Company Name",
    "Target Job Profiles",
    "Live Portal URL",
    "Hiring Status",
    "Last Scan Date",
    "Scanning Failure",
    "Needs Another Scan",
    "Notes",
    "Relevant Vacancies",
    "Actively Hiring Java/Spring",
    "Top Match Score",
    "ATS Checked",
    "Final URL",
]
_STATUS_RANK = {
    "Actively hiring Java/Spring": 0,
    "Priority roles found": 1,
    "Needs another scan": 2,
    "Scan failed": 3,
    "Not scanned": 4,
    "No priority roles on verified listing": 5,
}
_STATUS_FILL = {
    "Actively hiring Java/Spring": "C6EFCE",
    "Priority roles found": "DDEBF7",
    "Needs another scan": "FCE4D6",
    "Scan failed": "F8CBAD",
    "Not scanned": "FFF2CC",
    "No priority roles on verified listing": "E7E6E6",
}
_HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_WRAP = Alignment(wrap_text=True, vertical="top")


@dataclass
class ScanSummary:
    catalog_size: int
    scanned: int
    new_jobs: list[dict[str, object]]
    updated_jobs: int
    actively_hiring: list[str]
    needs_rescan: list[tuple[str, str, str]]
    output_path: Path


def vacancy_key(company: str, url: str, title: str, location: str) -> str:
    canon = canonical_url(url)
    if canon:
        return f"{company.casefold()}|url|{canon}"
    return f"{company.casefold()}|title|{title.casefold()}|{location.casefold()}"


def _text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _keeps_vacancy(row: dict[str, object], today: datetime) -> bool:
    return (
        experience_acceptable(_text(row.get("Experience"))) is True
        and posted_within_week(_text(row.get("Posting Date")), today.date()) is True
    )


def _score(value: object) -> int:
    try:
        return int(value)  # type: ignore[arg-type]
    except (TypeError, ValueError):
        return 0


def _excel_value(value: object) -> object:
    if isinstance(value, str) and value and value[0] in "=+-@":
        return "'" + value
    return value if value != "" else None


def _read_rows(sheet: Worksheet, headers: list[str]) -> list[dict[str, object]]:
    found = [_text(cell.value) for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    if found != headers:
        raise ValueError(
            f"Workbook sheet {sheet.title} is not the product Java scan format. "
            "Choose a different output path so an unrelated workbook is not overwritten."
        )
    rows: list[dict[str, object]] = []
    for cells in sheet.iter_rows(min_row=2, values_only=True):
        if not any(cell not in (None, "") for cell in cells):
            continue
        row = {header: cells[index] if index < len(cells) else "" for index, header in enumerate(headers)}
        if not _text(row.get("Company Name")):
            continue
        rows.append(row)
    return rows


def _job_row(job: FoundJob, *, active: str, first_seen: str, last_seen: str, status: str) -> dict[str, object]:
    return {
        "Company Name": job.company,
        "Job Title": job.title,
        "Location": job.location,
        "Experience": job.experience,
        "Required Skills": job.skills,
        "Job URL": job.url,
        "Posting Date": job.posting_date,
        "Match Score": job.match_score,
        "Role Match": job.role_match,
        "Verified Via": job.verified_via,
        "Location Check": job.location_check,
        "First Seen": first_seen,
        "Last Seen": last_seen,
        "Record Status": status,
        "Actively Hiring Java/Spring": active,
    }


def _same_vacancy(old: dict[str, object], new: dict[str, object]) -> bool:
    fields = [
        "Job Title",
        "Location",
        "Experience",
        "Required Skills",
        "Posting Date",
        "Role Match",
        "Verified Via",
        "Location Check",
        "Actively Hiring Java/Spring",
    ]
    return all(_text(old.get(field)) == _text(new.get(field)) for field in fields) and _score(
        old.get("Match Score")
    ) == _score(new.get("Match Score"))


def _style_header(sheet: Worksheet, width: dict[str, int]) -> None:
    header_row = next(sheet.iter_rows(min_row=1, max_row=1))
    for index, cell in enumerate(header_row, start=1):
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        name = str(cell.value or "")
        sheet.column_dimensions[get_column_letter(index)].width = width.get(name, 18)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 22


def _write_table(sheet: Worksheet, headers: list[str], rows: list[dict[str, object]], width: dict[str, int]) -> None:
    sheet.append(headers)
    for row in rows:
        sheet.append([_excel_value(row.get(header, "")) for header in headers])
    _style_header(sheet, width)
    wrap_headers = {"Required Skills", "Notes", "Scanning Failure", "Job URL", "Live Portal URL", "Final URL"}
    for column, header in enumerate(headers, start=1):
        if header not in wrap_headers:
            continue
        for cell in sheet.iter_rows(min_row=2, min_col=column, max_col=column):
            cell[0].alignment = _WRAP
    if "Hiring Status" in headers:
        status_col = headers.index("Hiring Status") + 1
        for row in sheet.iter_rows(min_row=2, min_col=status_col, max_col=status_col):
            fill = _STATUS_FILL.get(_text(row[0].value))
            if fill:
                row[0].fill = PatternFill("solid", fgColor=fill)


def _not_scanned(company: Company) -> dict[str, object]:
    return {
        "Company Name": company.name,
        "Target Job Profiles": company.target_profiles,
        "Live Portal URL": company.portal_url,
        "Hiring Status": "Not scanned",
        "Last Scan Date": "",
        "Scanning Failure": "",
        "Needs Another Scan": "Yes",
        "Notes": "Not included in this run.",
        "Relevant Vacancies": "",
        "Actively Hiring Java/Spring": "No",
        "Top Match Score": "",
        "ATS Checked": "",
        "Final URL": "",
    }


def update_workbook(
    path: Path,
    catalog: list[Company],
    results: list[CompanyResult],
    scanned_at: datetime,
) -> ScanSummary:
    path.parent.mkdir(parents=True, exist_ok=True)
    stamp = scanned_at.strftime("%Y-%m-%d %H:%M IST")
    if path.exists():
        workbook = load_workbook(path)
        old_jobs = _read_rows(workbook["Vacancies"], VACANCY_HEADERS) if "Vacancies" in workbook.sheetnames else None
        old_companies = (
            _read_rows(workbook["Company Status"], COMPANY_HEADERS)
            if "Company Status" in workbook.sheetnames
            else None
        )
        if old_jobs is None or old_companies is None:
            raise ValueError(
                f"{path} is not a product Java scan workbook. Choose a different output path."
            )
    else:
        workbook = Workbook()
        old_jobs = []
        old_companies = []

    old_job_by_key = {
        vacancy_key(
            _text(row.get("Company Name")),
            _text(row.get("Job URL")),
            _text(row.get("Job Title")),
            _text(row.get("Location")),
        ): row
        for row in old_jobs
    }
    old_company_by_name = {_text(row.get("Company Name")): row for row in old_companies}
    results_by_name = {result.company: result for result in results}
    new_keys: dict[str, set[str]] = {}
    for result in results:
        if not result.portal_loaded:
            continue
        keys = new_keys.setdefault(result.company, set())
        for job in result.jobs:
            keys.add(vacancy_key(job.company, job.url, job.title, job.location))

    merged: dict[str, dict[str, object]] = {}
    for key, old in old_job_by_key.items():
        if not _keeps_vacancy(old, scanned_at):
            continue
        company_name = _text(old.get("Company Name"))
        result = results_by_name.get(company_name)
        fresh_keys = new_keys.get(company_name, set())
        if result and result.portal_loaded and key in fresh_keys:
            continue
        row = {header: old.get(header, "") for header in VACANCY_HEADERS}
        if result and result.portal_loaded and result.listing_complete:
            row["Record Status"] = "Not seen on latest verified scan"
            row["Actively Hiring Java/Spring"] = "Yes" if result.actively_hiring else "No"
        elif _text(row.get("Record Status")) == "New":
            row["Record Status"] = "Carried forward"
        merged[key] = row

    new_rows: list[dict[str, object]] = []
    updated = 0
    for result in results:
        if not result.portal_loaded:
            continue
        active = "Yes" if result.actively_hiring else "No"
        for job in result.jobs:
            key = vacancy_key(job.company, job.url, job.title, job.location)
            old = old_job_by_key.get(key)
            if old is None:
                row = _job_row(job, active=active, first_seen=stamp, last_seen=stamp, status="New")
            else:
                row = _job_row(
                    job,
                    active=active,
                    first_seen=_text(old.get("First Seen")) or stamp,
                    last_seen=stamp,
                    status="Unchanged",
                )
                if _same_vacancy(old, row):
                    row["Record Status"] = "Unchanged"
                else:
                    row["Record Status"] = "Updated"
                    updated += 1
            if not _keeps_vacancy(row, scanned_at):
                continue
            if old is None:
                new_rows.append(row)
            merged[key] = row

    vacancy_rows = sorted(
        merged.values(),
        key=lambda row: (
            0 if _text(row.get("Actively Hiring Java/Spring")) == "Yes" else 1,
            -_score(row.get("Match Score")),
            _text(row.get("Company Name")).casefold(),
            _text(row.get("Job Title")).casefold(),
        ),
    )

    company_rows: list[dict[str, object]] = []
    for company in catalog:
        result = results_by_name.get(company.name)
        previous = old_company_by_name.get(company.name)
        if result is None:
            company_rows.append(previous or _not_scanned(company))
            continue
        kept_rows = [row for row in vacancy_rows if _text(row.get("Company Name")) == company.name]
        if result.portal_loaded:
            top_score: object = max((job.match_score for job in result.jobs), default="")
            relevant: object = len(result.jobs)
            active = "Yes" if result.actively_hiring else "No"
        elif kept_rows:
            top_score = max(_score(row.get("Match Score")) for row in kept_rows)
            relevant = len(kept_rows)
            active = "Yes" if any(_text(row.get("Actively Hiring Java/Spring")) == "Yes" for row in kept_rows) else "No"
        else:
            top_score = ""
            relevant = 0 if previous else ""
            active = "No"
        notes = result.notes
        if not result.portal_loaded and any(
            _text(row.get("Company Name")) == company.name for row in vacancy_rows
        ):
            notes = f"{notes} Previously saved vacancies were kept.".strip()
        company_rows.append(
            {
                "Company Name": company.name,
                "Target Job Profiles": company.target_profiles,
                "Live Portal URL": company.portal_url,
                "Hiring Status": result.hiring_status,
                "Last Scan Date": stamp,
                "Scanning Failure": result.scanning_failure,
                "Needs Another Scan": "Yes" if result.needs_rescan else "No",
                "Notes": notes,
                "Relevant Vacancies": relevant,
                "Actively Hiring Java/Spring": active or "No",
                "Top Match Score": top_score if top_score != "" else "",
                "ATS Checked": result.ats_checked,
                "Final URL": result.final_url,
            }
        )

    company_rows.sort(
        key=lambda row: (
            _STATUS_RANK.get(_text(row.get("Hiring Status")), 9),
            0 if _text(row.get("Actively Hiring Java/Spring")) == "Yes" else 1,
            -_score(row.get("Top Match Score")),
            _text(row.get("Company Name")).casefold(),
        )
    )

    workbook = Workbook()
    vacancies = workbook.active
    vacancies.title = "Vacancies"
    status_sheet = workbook.create_sheet("Company Status", 1)
    report = workbook.create_sheet("Scan Report", 2)
    _write_table(
        vacancies,
        VACANCY_HEADERS,
        vacancy_rows,
        {
            "Company Name": 28,
            "Job Title": 42,
            "Location": 28,
            "Experience": 18,
            "Required Skills": 42,
            "Job URL": 54,
            "Posting Date": 18,
            "Match Score": 14,
            "Role Match": 28,
            "Verified Via": 22,
            "Location Check": 18,
            "First Seen": 20,
            "Last Seen": 20,
            "Record Status": 32,
            "Actively Hiring Java/Spring": 28,
        },
    )
    _write_table(
        status_sheet,
        COMPANY_HEADERS,
        company_rows,
        {
            "Company Name": 28,
            "Target Job Profiles": 28,
            "Live Portal URL": 54,
            "Hiring Status": 38,
            "Last Scan Date": 20,
            "Scanning Failure": 36,
            "Needs Another Scan": 22,
            "Notes": 55,
            "Relevant Vacancies": 20,
            "Actively Hiring Java/Spring": 28,
            "Top Match Score": 18,
            "ATS Checked": 36,
            "Final URL": 54,
        },
    )
    _write_report(report, catalog, results, vacancy_rows, company_rows, updated, stamp)
    workbook.save(path)

    rescan = [
        (
            _text(row.get("Company Name")),
            _text(row.get("Live Portal URL")),
            _text(row.get("Scanning Failure")) or _text(row.get("Notes")) or _text(row.get("Hiring Status")),
        )
        for row in company_rows
        if _text(row.get("Needs Another Scan")) == "Yes"
    ]
    return ScanSummary(
        catalog_size=len(catalog),
        scanned=len(results),
        new_jobs=new_rows,
        updated_jobs=updated,
        actively_hiring=[
            _text(row.get("Company Name"))
            for row in company_rows
            if _text(row.get("Actively Hiring Java/Spring")) == "Yes"
            and _text(row.get("Hiring Status")) == "Actively hiring Java/Spring"
        ],
        needs_rescan=rescan,
        output_path=path,
    )


def _write_report(
    sheet: Worksheet,
    catalog: list[Company],
    results: list[CompanyResult],
    vacancy_rows: list[dict[str, object]],
    company_rows: list[dict[str, object]],
    updated: int,
    stamp: str,
) -> None:
    new_jobs = [row for row in vacancy_rows if _text(row.get("Record Status")) == "New"]
    rescan_rows = [row for row in company_rows if _text(row.get("Needs Another Scan")) == "Yes"]
    active_count = sum(
        1
        for row in company_rows
        if _text(row.get("Hiring Status")) == "Actively hiring Java/Spring"
    )
    sheet["A1"] = "Product company Java backend scan"
    sheet["A1"].font = Font(bold=True, size=14)
    summary = [
        ("Run time", stamp),
        ("Catalog companies", len(catalog)),
        ("Companies scanned this run", len(results)),
        ("Newly discovered matching jobs", len(new_jobs)),
        ("Updated jobs", updated),
        ("Companies actively hiring Java/Spring", active_count),
        ("Companies requiring another scan", len(rescan_rows)),
    ]
    for index, (label, value) in enumerate(summary, start=2):
        sheet.cell(index, 1, label).font = Font(bold=True)
        sheet.cell(index, 2, value)
    header_row = 10
    new_headers = ["Company Name", "Job Title", "Location", "Match Score", "Role Match", "Job URL"]
    sheet.cell(header_row, 1, "Newly discovered jobs").font = Font(bold=True, size=12)
    for column, header in enumerate(new_headers, start=1):
        cell = sheet.cell(header_row + 1, column, header)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
    for offset, row in enumerate(new_jobs, start=header_row + 2):
        for column, header in enumerate(new_headers, start=1):
            sheet.cell(offset, column, _excel_value(row.get(header, "")))
    rescan_header = header_row + 3 + max(len(new_jobs), 1)
    sheet.cell(rescan_header, 1, "Companies requiring another scan").font = Font(bold=True, size=12)
    rescan_headers = ["Company Name", "Live Portal URL", "Hiring Status", "Reason"]
    for column, header in enumerate(rescan_headers, start=1):
        cell = sheet.cell(rescan_header + 1, column, header)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
    for offset, row in enumerate(rescan_rows, start=rescan_header + 2):
        reason = _text(row.get("Scanning Failure")) or _text(row.get("Notes"))
        values = [
            row.get("Company Name"),
            row.get("Live Portal URL"),
            row.get("Hiring Status"),
            reason,
        ]
        for column, value in enumerate(values, start=1):
            sheet.cell(offset, column, _excel_value(value))
    sheet.column_dimensions["A"].width = 42
    sheet.column_dimensions["B"].width = 46
    sheet.column_dimensions["C"].width = 38
    sheet.column_dimensions["D"].width = 28
    sheet.column_dimensions["E"].width = 24
    sheet.column_dimensions["F"].width = 54
    sheet.freeze_panes = "A2"
