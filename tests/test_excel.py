"""Test Excel export."""

from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

from job_agent.export.excel_exporter import ExcelExporter
from job_agent.matching.scorer import MatchScorer
from job_agent.normalization.filters import normalize_job
from job_agent.schemas import RawJob


def test_excel_export(tmp_path, resume):
    scorer = MatchScorer(resume)

    raw = RawJob(
        source="test",
        title="Android Engineer",
        company="TestCorp",
        location="Mumbai",
        description="Kotlin, Jetpack Compose, MVVM, 2+ years",
        job_url="https://example.com/job",
        application_url="https://example.com/apply",
        date_posted=datetime.utcnow(),
    )
    job = normalize_job(raw)
    match = scorer.score(job)

    exporter = ExcelExporter()
    path = tmp_path / "test_jobs.xlsx"
    exporter.export([match], path)

    assert path.exists()
    wb = load_workbook(str(path))
    ws = wb.active
    assert ws.title.startswith("Top")
    assert ws.cell(row=1, column=1).value == "Rank"
    assert ws.cell(row=2, column=2).value == "Android Engineer"
    assert ws.cell(row=2, column=3).value == "TestCorp"
