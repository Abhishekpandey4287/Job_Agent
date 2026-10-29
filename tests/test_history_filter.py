"""History comparison against earlier job Excel files."""

from types import SimpleNamespace

from openpyxl import Workbook

from job_agent.export.history_filter import HistoryRules, exclude_known_jobs


def _job(**kwargs):
    defaults = dict(
        title="Android Engineer",
        company="Acme",
        job_url="https://boards.greenhouse.io/acme/jobs/99",
        application_url="https://boards.greenhouse.io/acme/jobs/99?utm=email",
        source_job_id="99",
        career_portal_url="",
        career_portal_job_id="",
    )
    defaults.update(kwargs)
    return SimpleNamespace(**defaults)


def _row(job):
    return (SimpleNamespace(job=job), "Android", "Android")


def _write_history(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.append(["Job Title", "Company", "Job ID", "Direct Application Link", "Status"])
    for title, company, job_id, url, status in rows:
        ws.append([title, company, job_id, url, status])
        if url:
            ws.cell(row=ws.max_row, column=4).hyperlink = url
            ws.cell(row=ws.max_row, column=4).value = "Apply Here"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def test_drops_jobs_already_listed_or_applied(tmp_path):
    previous = tmp_path / "2026-10-01" / "android_backend_jobs_2026-10-01.xlsx"
    _write_history(
        previous,
        [
            ("Android Engineer", "Acme", "99", "https://boards.greenhouse.io/acme/jobs/99", "DISCOVERED"),
            ("Backend Engineer", "Northwind", "55", "https://jobs.example.com/55", "Applied"),
        ],
    )
    rows = [
        _row(_job()),
        _row(_job(
            title="Backend Engineer",
            company="Northwind Labs",
            job_url="https://other.example/listing",
            application_url="https://jobs.example.com/55",
            source_job_id="",
        )),
        _row(_job(
            title="Kotlin Developer",
            company="Fresh Co",
            job_url="https://jobs.example.com/new",
            application_url="https://jobs.example.com/new",
            source_job_id="new-1",
        )),
    ]

    report = exclude_known_jobs(rows, tmp_path)

    assert report.removed_seen == 1
    assert report.removed_applied == 1
    assert len(report.kept) == 1
    assert report.kept[0][0].job.company == "Fresh Co"


def test_applied_status_still_drops_when_seen_jobs_are_kept(tmp_path):
    previous = tmp_path / "old.xlsx"
    _write_history(
        previous,
        [
            ("Android Engineer", "Acme", "", "https://jobs.example.com/seen", "DISCOVERED"),
            ("Backend Engineer", "Kept", "", "https://jobs.example.com/applied", "INTERVIEW"),
        ],
    )
    rows = [
        _row(_job(
            title="Android Engineer",
            company="Acme",
            job_url="https://jobs.example.com/seen",
            application_url="https://jobs.example.com/seen",
            source_job_id="",
        )),
        _row(_job(
            title="Backend Engineer",
            company="Kept",
            job_url="https://jobs.example.com/applied",
            application_url="https://jobs.example.com/applied",
            source_job_id="",
        )),
    ]
    rules = HistoryRules(drop_previously_seen=False)

    report = exclude_known_jobs(rows, tmp_path, rules=rules)

    assert report.removed_seen == 0
    assert report.removed_applied == 1
    assert [row[0].job.company for row in report.kept] == ["Acme"]


def test_same_day_file_is_not_used_as_its_own_history(tmp_path):
    current = tmp_path / "2026-10-06" / "android_backend_jobs_2026-10-06.xlsx"
    _write_history(
        current,
        [("Android Engineer", "Acme", "99", "https://jobs.example.com/99", "DISCOVERED")],
    )
    rows = [_row(_job(
        job_url="https://jobs.example.com/99",
        application_url="https://jobs.example.com/99",
    ))]

    report = exclude_known_jobs(rows, tmp_path, skip_file=current)

    assert report.workbooks == []
    assert report.kept == rows
