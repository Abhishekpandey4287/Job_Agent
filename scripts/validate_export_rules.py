#!/usr/bin/env python3
"""Post-run validator: Excel + top_matches against score/freshness/location rules."""

from __future__ import annotations

import json
import sys
from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from job_agent.matching.android_stack import is_excluded_cross_platform_primary, is_kotlin_android_role


def _find_today_output() -> Path:
    out = ROOT / "output"
    today = datetime.utcnow().strftime("%Y-%m-%d")
    d = out / today
    if d.exists():
        return d
    # newest dated folder
    dirs = sorted([p for p in out.iterdir() if p.is_dir()], reverse=True)
    return dirs[0] if dirs else d


def main() -> int:
    out_dir = _find_today_output()
    excel = next(out_dir.glob("*.xlsx"), None)
    snap = out_dir / "top_matches.json"
    print(f"Output dir: {out_dir}")
    print(f"Excel: {excel}")
    print(f"Snapshot: {snap.exists()}")

    errors: list[str] = []
    if not excel or not excel.exists():
        print("FAIL: no Excel")
        return 1
    if not snap.exists():
        print("FAIL: no top_matches.json")
        return 1

    wb = load_workbook(excel, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    headers = [str(h or "").strip() for h in rows[0]]
    col = {h: i for i, h in enumerate(headers)}
    print("Headers:", headers[:12], "...")

    def g(row, *names, default=None):
        for n in names:
            if n in col:
                return row[col[n]]
        return default

    data = rows[1:]
    scores = []
    ages_h = []
    unknown = 0
    india = 0
    remote = 0
    flutter_android = 0
    now = datetime.utcnow()

    for row in data:
        if not row or all(c is None for c in row):
            continue
        score = g(row, "Match Score", "Score")
        try:
            score_f = float(score)
        except (TypeError, ValueError):
            errors.append(f"bad score: {score}")
            continue
        scores.append(score_f)
        if score_f < 75:
            errors.append(f"score<{75}: {score_f} title={g(row,'Job Title','Title')}")

        posted = g(row, "Date Posted", "Posted Date", "date_posted")
        if posted is None or posted == "":
            unknown += 1
            errors.append(f"undated: {g(row,'Job Title','Title')}")
        else:
            if isinstance(posted, datetime):
                dt = posted.replace(tzinfo=None) if posted.tzinfo else posted
            else:
                try:
                    dt = datetime.fromisoformat(str(posted).replace("Z", ""))
                except Exception:
                    unknown += 1
                    errors.append(f"unparseable date {posted}")
                    continue
            age = (now - dt).total_seconds() / 3600.0
            ages_h.append(age)
            if age > 24.01:
                errors.append(f"age>{24}h ({age:.1f}h): {g(row,'Job Title','Title')}")

        loc = str(g(row, "Location", default="") or "")
        if "remote" in loc.lower():
            remote += 1
        else:
            india += 1

        title = str(g(row, "Job Title", "Title", default="") or "")
        profile = str(g(row, "Resume Used", "Profile", "Matched Profile", default="") or "")
        if profile.lower().startswith("android") or profile.lower() == "android":
            if is_excluded_cross_platform_primary(title) or not is_kotlin_android_role(title):
                flutter_android += 1
                errors.append(f"Android profile got non-Kotlin/Android primary: {title}")

    data_snap = json.loads(snap.read_text())
    # snapshot may be list or {matches: [...]}
    matches = data_snap if isinstance(data_snap, list) else data_snap.get("matches") or data_snap.get("jobs") or []
    snap_scores = []
    for m in matches:
        if isinstance(m, dict):
            sc = m.get("match_score") or m.get("score")
            if sc is not None:
                snap_scores.append(float(sc))
                if float(sc) < 75:
                    errors.append(f"snapshot score<75: {sc}")
            job = m.get("job") or m
            title = (job.get("title") or "") if isinstance(job, dict) else ""
            resume = m.get("resume_used") or m.get("profile") or ""
            if str(resume).lower().startswith("android") and title:
                if is_excluded_cross_platform_primary(title):
                    errors.append(f"snapshot Android Flutter/RN: {title}")

    print(f"Excel rows: {len(scores)}")
    print(f"min_score: {min(scores) if scores else None}")
    print(f"max_age_h: {max(ages_h) if ages_h else None}")
    print(f"unknown_dates: {unknown}")
    print(f"india-ish: {india} remote-ish: {remote}")
    print(f"snapshot matches: {len(matches)} min_score={min(snap_scores) if snap_scores else None}")
    print(f"android_stack_violations: {flutter_android}")

    if errors:
        print(f"\nFAIL ({len(errors)} issues):")
        for e in errors[:30]:
            print(" -", e)
        return 2
    print("\nPASS: export rules satisfied")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
