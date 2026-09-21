#!/usr/bin/env python3
"""Migrate resume folders/files to company-based naming convention."""

import sqlite3
import sys
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

from openpyxl import load_workbook

from job_agent.resume.folder_names import (
    basenames_for_company,
    company_folder_name,
    format_role_label,
    new_company_basename_tracker,
    resume_file_basename,
)


def main():
    date = sys.argv[1] if len(sys.argv) > 1 else "2026-09-02"
    resumes_root = PROJECT_ROOT / "output" / date / "resumes"
    excel_path = PROJECT_ROOT / "output" / date / f"android_backend_jobs_{date}.xlsx"

    if not resumes_root.exists():
        print(f"No resumes folder: {resumes_root}")
        return 1

    tracker = new_company_basename_tracker()

    if excel_path.exists():
        wb = load_workbook(str(excel_path))
        ws = wb.active
        for row in range(2, ws.max_row + 1):
            company = ws.cell(row=row, column=4).value
            title = ws.cell(row=row, column=3).value
            role_cat = ws.cell(row=row, column=2).value or "Android"
            if not company:
                continue
            folder = company_folder_name(company)
            target_dir = resumes_root / folder
            target_dir.mkdir(parents=True, exist_ok=True)
            file_base = resume_file_basename(
                role_cat, title or "", basenames_for_company(tracker, company)
            )
            # Find any existing tex in old company-like dirs
            for old_dir in resumes_root.iterdir():
                if not old_dir.is_dir() or old_dir.name in ("master", folder):
                    continue
                if company.lower().replace(" ", "")[:8] not in old_dir.name.lower().replace("_", ""):
                    continue
                for tex in old_dir.glob("*.tex"):
                    new_tex = target_dir / f"{file_base}.tex"
                    tex.rename(new_tex)
                    pdf = tex.with_suffix(".pdf")
                    if pdf.exists():
                        pdf.rename(target_dir / f"{file_base}.pdf")
                    print(f"Moved {tex.name} -> {folder}/{file_base}.tex")
            pdf = target_dir / f"{file_base}.pdf"
            tex = target_dir / f"{file_base}.tex"
            if pdf.exists():
                c = ws.cell(row=row, column=18)
                c.value = pdf.name
                c.hyperlink = f"file://{pdf.resolve()}"
            if tex.exists():
                c = ws.cell(row=row, column=19)
                c.value = tex.name
                c.hyperlink = f"file://{tex.resolve()}"
        wb.save(str(excel_path))
        print(f"Updated Excel: {excel_path}")

    print("\nCompany folders:")
    for d in sorted(resumes_root.iterdir()):
        if d.is_dir() and d.name != "master":
            files = list(d.glob("Abhishek-Pandey_*"))
            print(f"  {d.name}/ ({len(files)} resume files)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
