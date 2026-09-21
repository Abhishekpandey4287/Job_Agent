#!/usr/bin/env python3
"""Re-score an existing results Excel with the current scoring engine.

Rewrites only the score-derived columns (Role Category, Match Score, Match
Category, Resume Used, skills, explanation) and re-ranks rows. Hyperlinks,
generated resume/cover paths, Status and Notes are carried across untouched.

Usage:
  .venv/bin/python3 scripts/rescore_excel.py                    # latest output day
  .venv/bin/python3 scripts/rescore_excel.py output/2026-09-15
  .venv/bin/python3 scripts/rescore_excel.py --dry-run
"""

from __future__ import annotations

import json
import os
import shutil
import sys
from copy import copy
from datetime import datetime
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parent.parent
os.chdir(PROJECT_ROOT)
sys.path.insert(0, str(PROJECT_ROOT))

VENV_PYTHON = PROJECT_ROOT / ".venv" / "bin" / "python3"
if VENV_PYTHON.exists() and Path(sys.executable).resolve() != VENV_PYTHON.resolve():
    os.execv(str(VENV_PYTHON), [str(VENV_PYTHON), *sys.argv])

from openpyxl import load_workbook  # noqa: E402

from job_agent.matching.multi_profile import score_job_multi  # noqa: E402
from job_agent.profiles import load_profiles  # noqa: E402
from job_agent.resume.folder_names import format_role_label  # noqa: E402
from job_agent.schemas import NormalizedJob  # noqa: E402

SCORE_COLUMNS = (
    "Role Category",
    "Match Score",
    "Match Category",
    "Resume Used",
    "Matching Skills",
    "Missing Skills",
    "Match Explanation",
)


def resolve_day(arg: str | None) -> Path:
    if arg:
        path = Path(arg).expanduser().resolve()
        if not path.exists():
            raise SystemExit(f"Not found: {path}")
        return path
    days = sorted((PROJECT_ROOT / "output").glob("20*-*-*"), reverse=True)
    if not days:
        raise SystemExit("No output/<date> folders found")
    return days[0]


def resolve_excel(day: Path) -> Path:
    books = sorted(day.glob("*.xlsx"))
    if not books:
        raise SystemExit(f"No .xlsx under {day}")
    return books[0]


def key_of(company: str, title: str) -> tuple[str, str]:
    return (company or "").strip().lower(), (title or "").strip().lower()


def main() -> None:
    args = [a for a in sys.argv[1:] if a != "--dry-run"]
    dry_run = "--dry-run" in sys.argv

    day = resolve_day(args[0] if args else None)
    excel_path = resolve_excel(day)
    snapshot_path = day / "top_matches.json"
    if not snapshot_path.exists():
        raise SystemExit(f"Missing snapshot: {snapshot_path}")

    raw = json.loads(snapshot_path.read_text(encoding="utf-8"))
    items = raw if isinstance(raw, list) else raw.get("matches") or raw.get("items") or []
    by_key = {key_of(i["match"]["job"].get("company"), i["match"]["job"].get("title")): i for i in items}

    profiles = load_profiles()
    if not profiles:
        raise SystemExit("No resume profiles found")

    wb = load_workbook(excel_path)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    col = {h: n for n, h in enumerate(headers, 1)}
    for needed in (*SCORE_COLUMNS, "Rank", "Company", "Job Title"):
        if needed not in col:
            raise SystemExit(f"Excel missing column: {needed}")

    # Snapshot every row (values + styles + hyperlinks) so re-ranking can move whole rows.
    captured = []
    for r in range(2, ws.max_row + 1):
        cells = {}
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=r, column=c)
            cells[c] = {
                "value": cell.value,
                "font": copy(cell.font),
                "fill": copy(cell.fill),
                "alignment": copy(cell.alignment),
                "border": copy(cell.border),
                "number_format": cell.number_format,
                "hyperlink": cell.hyperlink.target if cell.hyperlink else None,
            }
        captured.append(cells)

    updated = 0
    unscored = []
    for cells in captured:
        company = str(cells[col["Company"]]["value"] or "")
        title = str(cells[col["Job Title"]]["value"] or "")
        item = by_key.get(key_of(company, title))
        if not item:
            unscored.append(f"{company} — {title} (no snapshot entry)")
            continue

        job = NormalizedJob(**item["match"]["job"])
        result = score_job_multi(job, profiles)
        if not result:
            unscored.append(f"{company} — {title} (filtered out by matcher)")
            continue

        match, track_label, resume_used = result
        cells[col["Role Category"]]["value"] = format_role_label(track_label)
        cells[col["Match Score"]]["value"] = match.match_score
        cells[col["Match Category"]]["value"] = match.match_category.value
        cells[col["Resume Used"]]["value"] = resume_used
        cells[col["Matching Skills"]]["value"] = ", ".join(match.matching_skills[:10])
        cells[col["Missing Skills"]]["value"] = ", ".join(match.missing_requirements[:8])
        cells[col["Match Explanation"]]["value"] = match.match_explanation
        cells["_match"] = match
        cells["_track"] = track_label
        cells["_resume_used"] = resume_used
        cells["_job"] = item
        updated += 1

    captured.sort(key=lambda c: float(c[col["Match Score"]]["value"] or 0), reverse=True)

    print(f"Excel:    {excel_path}")
    print(f"Snapshot: {snapshot_path}")
    print(f"Re-scored {updated}/{len(captured)} rows\n")
    print(f"{'Rank':>4}  {'Score':>6}  {'Resume':<10}  Title")
    for n, cells in enumerate(captured, 1):
        print(
            f"{n:>4}  {float(cells[col['Match Score']]['value'] or 0):>6.1f}  "
            f"{str(cells[col['Resume Used']]['value'] or ''):<10}  "
            f"{str(cells[col['Job Title']]['value'] or '')[:52]}"
        )
    if unscored:
        print("\nLeft unchanged:")
        for u in unscored:
            print(f"  - {u}")

    if dry_run:
        print("\n--dry-run: nothing written")
        return

    backup = excel_path.with_name(
        f"{excel_path.stem}.bak-{datetime.now():%Y%m%d-%H%M%S}{excel_path.suffix}"
    )
    shutil.copy2(excel_path, backup)

    for n, cells in enumerate(captured, 1):
        row = n + 1
        for c in range(1, len(headers) + 1):
            src = cells[c]
            cell = ws.cell(row=row, column=c)
            cell.value = src["value"]
            cell.font = src["font"]
            cell.fill = src["fill"]
            cell.alignment = src["alignment"]
            cell.border = src["border"]
            cell.number_format = src["number_format"]
            cell.hyperlink = src["hyperlink"]
        ws.cell(row=row, column=col["Rank"]).value = n
    wb.save(excel_path)

    # Keep the snapshot consistent so the resume agent starts from these scores.
    rescored_items = []
    for cells in captured:
        item = cells.get("_job")
        if not item:
            continue
        match = cells["_match"]
        item = json.loads(json.dumps(item, default=str))
        item["match"].update(json.loads(match.model_dump_json()))
        item["role_track"] = cells["_track"]
        item["resume_used"] = cells["_resume_used"]
        rescored_items.append(item)
    if rescored_items:
        snap_backup = snapshot_path.with_suffix(
            f".bak-{datetime.now():%Y%m%d-%H%M%S}.json"
        )
        shutil.copy2(snapshot_path, snap_backup)
        payload = rescored_items if isinstance(raw, list) else {**raw, "matches": rescored_items}
        snapshot_path.write_text(
            json.dumps(payload, indent=2, default=str), encoding="utf-8"
        )

    print(f"\nWrote {excel_path}")
    print(f"Backup {backup}")


if __name__ == "__main__":
    main()
