#!/usr/bin/env python3
"""Rewrite the Job ID column of an existing results Excel.

Keeps only the employer's own requisition/reference id. Job-board listing numbers
(Shine, Indeed, Naukri, …) and stray tracking URLs are cleared, because quoting a
board's listing number to a recruiter is worse than quoting nothing.

Usage:
  .venv/bin/python3 scripts/fix_job_ids.py                    # latest output day
  .venv/bin/python3 scripts/fix_job_ids.py output/2026-09-15
  .venv/bin/python3 scripts/fix_job_ids.py --dry-run
"""

from __future__ import annotations

import json
import os
import shutil
import sys
from datetime import datetime
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parent.parent
os.chdir(PROJECT_ROOT)
sys.path.insert(0, str(PROJECT_ROOT))

VENV_PYTHON = PROJECT_ROOT / ".venv" / "bin" / "python3"
if VENV_PYTHON.exists() and Path(sys.executable).resolve() != VENV_PYTHON.resolve():
    os.execv(str(VENV_PYTHON), [str(VENV_PYTHON), *sys.argv])

from openpyxl import load_workbook  # noqa: E402

from job_agent.discovery.job_ids import extract_company_reference_id  # noqa: E402


def resolve_day(arg: str | None) -> Path:
    if arg:
        path = Path(arg).expanduser().resolve()
        if not path.exists():
            raise SystemExit(f"Not found: {path}")
        return path
    days = sorted((PROJECT_ROOT / "output").glob("20*-*-*"), reverse=True)
    if not days:
        raise SystemExit("No output/<date> folders found")
    return days[0]


def key_of(company: str, title: str) -> tuple[str, str]:
    return (company or "").strip().lower(), (title or "").strip().lower()


def pick_job(candidates: list[dict], current_id: str) -> tuple[dict | None, str]:
    """Choose the snapshot job a spreadsheet row refers to.

    Returns (job, problem). Ambiguity is reported rather than guessed — writing the
    wrong requisition id is worse than leaving the cell alone.
    """
    if not candidates:
        return None, "no snapshot entry"
    if len(candidates) == 1:
        return candidates[0], ""

    current = (current_id or "").strip()
    if current:
        matched = [
            job
            for job in candidates
            if current == str(job.get("source_job_id") or "").strip()
            or current in str(job.get("job_url") or "")
            or current in str(job.get("application_url") or "")
        ]
        if len(matched) == 1:
            return matched[0], ""

    return None, f"ambiguous — {len(candidates)} postings share this company and title"


def main() -> None:
    args = [a for a in sys.argv[1:] if a != "--dry-run"]
    dry_run = "--dry-run" in sys.argv

    day = resolve_day(args[0] if args else None)
    books = [b for b in sorted(day.glob("*.xlsx")) if ".bak-" not in b.name]
    if not books:
        raise SystemExit(f"No .xlsx under {day}")
    excel_path = books[0]

    snapshot_path = day / "top_matches.json"
    if not snapshot_path.exists():
        raise SystemExit(f"Missing snapshot: {snapshot_path}")
    raw = json.loads(snapshot_path.read_text(encoding="utf-8"))
    items = raw if isinstance(raw, list) else raw.get("matches") or raw.get("items") or []

    # Company+title is not unique — one employer can post several distinct
    # requisitions under an identical title — so keep every candidate and
    # disambiguate below instead of letting the last one win.
    by_key: dict[tuple[str, str], list[dict]] = {}
    for item in items:
        job = item["match"]["job"]
        by_key.setdefault(key_of(job.get("company"), job.get("title")), []).append(job)

    wb = load_workbook(excel_path)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    col = {h: n for n, h in enumerate(headers, 1)}
    for needed in ("Job ID", "Company", "Job Title", "Job Source"):
        if needed not in col:
            raise SystemExit(f"Excel missing column: {needed}")

    kept = cleared = unknown = 0
    rows = []
    for r in range(2, ws.max_row + 1):
        company = str(ws.cell(r, col["Company"]).value or "")
        title = str(ws.cell(r, col["Job Title"]).value or "")
        before = str(ws.cell(r, col["Job ID"]).value or "")
        job, problem = pick_job(by_key.get(key_of(company, title), []), before)

        if job is None:
            unknown += 1
            rows.append((company, title, before, before, f"LEFT AS-IS — {problem}"))
            continue

        after = extract_company_reference_id(
            source=str(job.get("source") or ""),
            source_job_id=str(job.get("source_job_id") or ""),
            job_url=str(job.get("job_url") or ""),
            application_url=str(job.get("application_url") or ""),
        )
        if after:
            kept += 1
            note = "company ATS id"
        else:
            cleared += 1
            note = f"cleared ({job.get('source')} listing id)"
        rows.append((company, title, before, after, note))
        if not dry_run:
            ws.cell(r, col["Job ID"]).value = after or None

    print(f"Excel: {excel_path}\n")
    for company, title, before, after, note in rows:
        shown_after = after or "(blank)"
        print(f"  {company[:22]:<24} {before[:26]:<28} -> {shown_after[:26]:<28} {note}")
    print(f"\nkept={kept} cleared={cleared} unmatched={unknown}")

    if dry_run:
        print("\n--dry-run: nothing written")
        return

    backup = excel_path.with_name(
        f"{excel_path.stem}.bak-{datetime.now():%Y%m%d-%H%M%S}{excel_path.suffix}"
    )
    shutil.copy2(excel_path, backup)
    wb.save(excel_path)
    print(f"\nWrote  {excel_path}")
    print(f"Backup {backup}")


if __name__ == "__main__":
    main()
