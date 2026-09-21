"""Excel report generator with clickable links."""

from datetime import datetime
from pathlib import Path
from typing import List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from job_agent.discovery.job_ids import extract_company_reference_id
from job_agent.logging_config import get_logger
from job_agent.schemas import JobMatch

logger = get_logger(__name__)

HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
LINK_FONT = Font(color="0563C1", underline="single")


COLUMNS = [
    ("Rank", 6),
    ("Job Title", 35),
    ("Company", 25),
    ("Job ID", 22),
    ("Location", 20),
    ("Date Posted", 14),
    ("Employment Type", 14),
    ("Seniority Level", 14),
    ("Salary", 18),
    ("Match Score", 12),
    ("Match Category", 16),
    ("Key Matching Skills", 35),
    ("Missing/Weak Requirements", 35),
    ("Match Explanation", 50),
    ("Job Source", 12),
    ("Direct Application Link", 40),
    ("Customized Resume", 30),
    ("Resume .tex File", 30),
    ("Application Status", 18),
    ("Notes", 25),
    ("Career Portal", 18),
    ("Career Portal Job ID", 22),
    ("Career Portal URL", 40),
    ("Career Job ID Type", 18),
]


class ExcelExporter:
    def export(
        self,
        matches: List[JobMatch],
        output_path: Path,
        resume_paths: dict[str, tuple[str, str]] | None = None,
    ) -> Path:
        wb = Workbook()
        ws = wb.active
        ws.title = f"Top {len(matches)} Jobs"

        for col_idx, (header, width) in enumerate(COLUMNS, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(col_idx)].width = width

        ws.freeze_panes = "A2"

        resume_paths = resume_paths or {}

        for rank, match in enumerate(matches, 1):
            row = rank + 1
            job = match.job
            paths = resume_paths.get(job.job_id, ("", ""))

            date_posted = ""
            if job.date_posted:
                date_posted = job.date_posted.strftime("%Y-%m-%d")
                if job.date_posted_verified:
                    date_posted += f" ({job.date_posted_verified.value})"

            company_ref_id = extract_company_reference_id(
                source=job.source,
                source_job_id=job.source_job_id,
                job_url=job.job_url,
                application_url=job.application_url,
            )

            values = [
                rank,
                job.title,
                job.company,
                company_ref_id,
                job.location,
                date_posted,
                job.employment_type.value,
                job.seniority.value,
                job.salary or "N/A",
                match.match_score,
                match.match_category.value,
                ", ".join(match.matching_skills[:10]),
                ", ".join(match.missing_requirements[:8]),
                match.match_explanation,
                job.source,
                job.application_url or job.job_url,
                paths[0] if paths[0] else "Not generated",
                paths[1] if paths[1] else "Not generated",
                "DISCOVERED",
                "",
                getattr(job, "career_portal", "") or "",
                getattr(job, "career_portal_job_id", "") or "",
                getattr(job, "career_portal_url", "") or "",
                getattr(job, "career_job_id_type", "") or "",
            ]

            for col_idx, value in enumerate(values, 1):
                cell = ws.cell(row=row, column=col_idx, value=value)
                cell.alignment = Alignment(vertical="top", wrap_text=True)

                if col_idx == 16 and value and str(value).startswith("http"):
                    cell.hyperlink = value
                    cell.font = LINK_FONT
                    cell.value = "Apply Here"

                if col_idx == 17 and paths[0]:
                    cell.hyperlink = f"file://{paths[0]}"
                    cell.font = LINK_FONT
                    cell.value = Path(paths[0]).parent.name

                if col_idx == 18 and paths[1]:
                    cell.hyperlink = f"file://{paths[1]}"
                    cell.font = LINK_FONT
                    cell.value = Path(paths[1]).parent.name

                if col_idx == len(COLUMNS) - 1 and str(value).startswith("http"):
                    cell.hyperlink = str(value)
                    cell.font = LINK_FONT

        output_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(str(output_path))
        logger.info("excel_exported", path=str(output_path), rows=len(matches))
        return output_path

    def export_dual_track(
        self,
        rows: list,
        output_path: Path,
        resume_paths: dict | None = None,
        outreach_map: dict | None = None,
    ) -> Path:
        """Export with Role Category and Resume Used columns."""
        dual_columns = [
            ("Rank", 6), ("Role Category", 14), ("Job Title", 35), ("Company", 25),
            ("Job ID", 22), ("Location", 20), ("Remote", 10), ("Date Posted", 16),
            ("Employment Type", 14), ("Seniority Level", 14), ("Salary", 18),
            ("Match Score", 12), ("Match Category", 16), ("Resume Used", 14),
            ("Matching Skills", 35), ("Missing Skills", 35), ("Match Explanation", 50),
            ("Job Source", 14), ("All Sources", 28), ("Direct Application Link", 40),
            ("Company URL", 28), ("Customized Resume", 30), ("Resume .tex File", 30),
            ("Cover Letter", 30), ("LinkedIn Contacts", 35), ("Connection Notes", 30),
            ("Outreach Status", 18), ("Status", 14), ("Notes", 25),
            ("Career Portal", 18), ("Career Portal Job ID", 22),
            ("Career Portal URL", 40), ("Career Job ID Type", 18),
        ]
        career_url_col = len(dual_columns) - 1
        wb = Workbook()
        ws = wb.active
        ws.title = f"Top {len(rows)} Jobs"
        resume_paths = resume_paths or {}
        outreach_map = outreach_map or {}

        for col_idx, (header, width) in enumerate(dual_columns, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            ws.column_dimensions[get_column_letter(col_idx)].width = width
        ws.freeze_panes = "A2"

        for rank, (match, role_cat, resume_used) in enumerate(rows, 1):
            job = match.job
            paths = resume_paths.get(job.job_id, ("", "", ""))
            pdf_p, tex_p, cover_p = paths[0], paths[1], paths[2]
            outreach = outreach_map.get(job.job_id, {})
            contact_names = outreach.get("contacts", "")
            note_path = outreach.get("note_path", "")
            outreach_status = outreach.get("status", "")

            date_posted = "POSTING_DATE_UNVERIFIED"
            if job.date_posted:
                date_posted = job.date_posted.strftime("%Y-%m-%d %H:%M")
                if job.date_posted_verified and job.date_posted_verified.value == "verified":
                    date_posted += " (verified)"
                else:
                    date_posted += " (estimated)"

            pdf_link_label = Path(pdf_p).name if pdf_p else "Not generated"
            tex_link_label = Path(tex_p).name if tex_p else "Not generated"
            cover_link_label = Path(cover_p).name if cover_p else "Not generated"

            # Blank unless it is genuinely the employer's reference id — never a
            # portal listing number and never a raw URL.
            company_ref_id = extract_company_reference_id(
                source=job.source,
                source_job_id=job.source_job_id,
                job_url=job.job_url,
                application_url=job.application_url,
            )

            all_sources = job.source
            reason = getattr(job, "relevance_reason", "") or ""
            if reason.startswith("[sources:"):
                end = reason.find("]")
                if end > 0:
                    all_sources = reason[len("[sources:"):end].strip()

            is_remote = "Yes" if (
                getattr(job, "is_remote", False)
                or "remote" in (job.location or "").lower()
            ) else "No"
            company_url = getattr(job, "company_url", "") or ""

            values = [
                rank, role_cat, job.title, job.company, company_ref_id, job.location,
                is_remote, date_posted,
                job.employment_type.value, job.seniority.value, job.salary or "N/A",
                match.match_score, match.match_category.value, resume_used,
                ", ".join(match.matching_skills[:10]),
                ", ".join(match.missing_requirements[:8]),
                match.match_explanation, job.source, all_sources,
                job.application_url or job.job_url,
                company_url or "—",
                pdf_link_label,
                tex_link_label,
                cover_link_label,
                contact_names or "—",
                Path(note_path).name if note_path else "—",
                outreach_status or "—",
                "DISCOVERED", "",
                getattr(job, "career_portal", "") or "",
                getattr(job, "career_portal_job_id", "") or "",
                getattr(job, "career_portal_url", "") or "",
                getattr(job, "career_job_id_type", "") or "",
            ]
            for col_idx, value in enumerate(values, 1):
                cell = ws.cell(row=rank + 1, column=col_idx, value=value)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if col_idx == 20 and str(value).startswith("http"):
                    cell.hyperlink = str(value)
                    cell.font = LINK_FONT
                    cell.value = "Apply Here"
                if col_idx == 21 and str(value).startswith("http"):
                    cell.hyperlink = str(value)
                    cell.font = LINK_FONT
                if col_idx == 22 and pdf_p and Path(pdf_p).exists():
                    cell.hyperlink = f"file://{pdf_p}"
                    cell.font = LINK_FONT
                if col_idx == 23 and tex_p and Path(tex_p).exists():
                    cell.hyperlink = f"file://{tex_p}"
                    cell.font = LINK_FONT
                if col_idx == 24 and cover_p and Path(cover_p).exists():
                    cell.hyperlink = f"file://{cover_p}"
                    cell.font = LINK_FONT
                if col_idx == 26 and note_path and Path(note_path).exists():
                    cell.hyperlink = f"file://{note_path}"
                    cell.font = LINK_FONT
                if col_idx == career_url_col and str(value).startswith("http"):
                    cell.hyperlink = str(value)
                    cell.font = LINK_FONT

        output_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(str(output_path))
        logger.info("excel_dual_exported", path=str(output_path), rows=len(rows))
        return output_path
